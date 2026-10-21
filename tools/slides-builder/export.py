"""Render Slides-format decks to images, then build PDF and PPTX.

usage: python3 export.py <deck_root> <poster.jpg> <video.mp4> <out_basename> <workdir>
"""
import asyncio, base64, html, json, os, re, sys
from playwright.async_api import async_playwright
from preview import FONT_CSS

ICON_DIR = "/home/claude/corum/icons/node_modules/lucide-static/icons"
ICON_MAP = {"Activity": "activity", "Book": "book-open", "Chart": "chart-column", "Chat": "message-square", "Check": "check",
            "CheckCircle": "circle-check", "Clock": "clock", "Cloud": "cloud", "Code": "code", "Database": "database", "Globe": "globe",
            "GraduationCap": "graduation-cap", "Home": "house", "Key": "key", "Lightbulb": "lightbulb", "Lightning": "zap", "Link": "link",
            "Lock": "lock", "PaperPlane": "send", "Play": "play", "Search": "search", "Settings": "settings", "Star": "star",
            "ThumbsUp": "thumbs-up", "Tool": "wrench", "Trust": "shield-check", "Users": "users", "Verified": "badge-check",
            "Warning": "triangle-alert", "Wrench": "wrench"}


def icon_svg(name, style):
    fn = ICON_MAP.get(name, "circle")
    p = os.path.join(ICON_DIR, fn + ".svg")
    if not os.path.exists(p):
        p = os.path.join(ICON_DIR, "bar-chart-3.svg") if name == "Chart" else os.path.join(ICON_DIR, "circle.svg")
    s = open(p).read()
    s = re.sub(r"<!--.*?-->", "", s, flags=re.S).strip()
    s = re.sub(r'\s(class|width|height)="[^"]*"', "", s, count=3)
    s = s.replace("<svg", f'<svg style="{style};flex:none;display:block"', 1)
    return s


def convert(sec, poster):
    sec = re.sub(r"<aside>.*?</aside>", "", sec, flags=re.S)

    def emb(m):
        st, inner = m.group(1), m.group(2)
        return f'<iframe style="{st};border:0" srcdoc="{html.escape(inner, quote=True)}"></iframe>'
    sec = re.sub(r'<x-embed[^>]*style="([^"]*)"[^>]*>(.*?)</x-embed>', emb, sec, flags=re.S)

    def svg2img(m):
        s = m.group(0)
        mm = re.search(r'style="([^"]*)"', s)
        st = mm.group(1) if mm else ""
        if "xmlns=" not in s[:200]:
            s = s.replace("<svg", '<svg xmlns="http://www.w3.org/2000/svg"', 1)
        s2 = s.replace(f'style="{st}"', "", 1)
        return f'<img src="data:image/svg+xml;base64,{base64.b64encode(s2.encode()).decode()}" style="{st}">'
    sec = re.sub(r"<svg[ >].*?</svg>", svg2img, sec, flags=re.S)
    # icons after svg conversion so their inline svg stays inline (currentColor needs the DOM)
    sec = re.sub(r'<x-icon name="([^"]*)" style="([^"]*)"></x-icon>', lambda m: icon_svg(m.group(1), m.group(2)), sec)
    sec = re.sub(r'src="/_blob/[0-9a-f]+"', f'src="file://{poster}"', sec)
    sec = re.sub(r'<section ([^>]*)style="', r'<div class="sl" \1style="position:relative;width:1920px;height:1080px;box-sizing:border-box;overflow:hidden;', sec, count=1)
    sec = re.sub(r"</section>\s*$", "</div>", sec.strip())
    return sec


def notes_of(sec):
    m = re.search(r"<aside>(.*?)</aside>", sec, flags=re.S)
    return html.unescape(re.sub(r"<[^>]+>", "", m.group(1))).strip() if m else ""


def title_of(sec):
    for tag in ("h1", "h2"):
        m = re.search(rf"<{tag}[^>]*>(.*?)</{tag}>", sec, flags=re.S)
        if m:
            return html.unescape(re.sub(r"<[^>]+>", "", m.group(1))).strip()
    return ""


async def render(root, poster, work):
    order = json.load(open(os.path.join(root, "project/deck.json")))["order"]
    os.makedirs(work, exist_ok=True)
    meta = []
    async with async_playwright() as p:
        b = await p.chromium.launch(args=["--allow-file-access-from-files"])
        pg = await b.new_page(viewport={"width": 1920, "height": 1080}, device_scale_factor=1.5)
        for i, sid in enumerate(order):
            sec = open(os.path.join(root, f"project/slides/{sid}.html")).read()
            doc = (f"<!doctype html><html><head><meta charset='utf-8'><style>{FONT_CSS} body{{margin:0}} p,h1,h2,h3,ul,ol{{margin:0}} "
                   f"ul,ol{{padding-left:1.2em}} .sl{{font-size:32px}} .sl *{{box-sizing:border-box}}</style></head><body>{convert(sec, poster)}</body></html>")
            fp = os.path.join(work, f"_{sid}.html")
            open(fp, "w").write(doc)
            await pg.goto("file://" + fp)
            await pg.wait_for_timeout(2600)
            out = os.path.join(work, f"{i + 1:02d}_{sid}.jpg")
            await pg.screenshot(path=out, type="jpeg", quality=90, clip={"x": 0, "y": 0, "width": 1920, "height": 1080})
            meta.append({"sid": sid, "img": out, "notes": notes_of(sec), "title": title_of(sec), "video": 'data-video="' in sec})
        await b.close()
    return meta


def build_pdf(meta, path):
    import img2pdf
    layout = img2pdf.get_layout_fun((img2pdf.in_to_pt(13.333), img2pdf.in_to_pt(7.5)))
    with open(path, "wb") as f:
        f.write(img2pdf.convert([m["img"] for m in meta], layout_fun=layout))


def build_pptx(meta, path, video, poster, deck_title):
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    blank = prs.slide_layouts[6]
    prs.core_properties.title = deck_title
    prs.core_properties.author = "Prof. Dr. Ahmet Altun"
    for m in meta:
        s = prs.slides.add_slide(blank)
        if m["video"]:
            mv = s.shapes.add_movie(video, 0, 0, prs.slide_width, prs.slide_height, poster_frame_image=poster, mime_type="video/mp4")
            mv.name = "Opening film"
        else:
            pic = s.shapes.add_picture(m["img"], 0, 0, prs.slide_width, prs.slide_height)
            pic.name = m["sid"]
            pic._element.nvPicPr.cNvPr.set("descr", m["title"] or m["sid"])
        if m["notes"]:
            s.notes_slide.notes_text_frame.text = m["notes"]
    prs.save(path)


if __name__ == "__main__":
    root, poster, video, outbase, work = sys.argv[1:6]
    meta = asyncio.run(render(root, poster, work))
    title = json.load(open(os.path.join(root, "project/deck.json")))["title"]
    build_pdf(meta, outbase + ".pdf")
    build_pptx(meta, outbase + ".pptx", video, poster, title)
    print("ok", len(meta), outbase)
